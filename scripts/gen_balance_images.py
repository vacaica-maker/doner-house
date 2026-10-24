import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

SRC = "Estados_Financieros_Doner_Street_2026.xlsx"
OUT_DIR = "imagenes/graficas"

BROWN = "#84592B"
DARK = "#442D1C"
LIGHT = "#F6EBD6"
WHITE = "#FFFDF8"
RED = "#B3261E"
GREEN = "#2E7D4F"

GROUP_LABELS = {
    "INGRESOS OPERACIONALES", "COSTO DE VENTAS",
    "GASTOS OPERACIONALES DE ADMINISTRACIÓN", "GASTOS OPERACIONALES DE VENTAS",
    "PARTIDAS NO OPERACIONALES", "ACTIVIDADES DE OPERACIÓN",
    "ACTIVIDADES DE INVERSIÓN", "ACTIVIDADES DE FINANCIACIÓN",
    "ACTIVO CORRIENTE", "ACTIVO NO CORRIENTE", "PASIVO CORRIENTE", "PATRIMONIO",
}

TOTAL_LABELS = {
    "Ventas netas", "Total costo de ventas", "Utilidad bruta",
    "Total gastos de administración", "Total gastos de ventas",
    "Utilidad operacional (EBIT)", "Utilidad antes de impuestos", "Utilidad neta",
    "Efectivo neto de operación", "Flujo neto del periodo", "Saldo final de efectivo",
    "Total activo corriente", "Total activo no corriente", "TOTAL ACTIVO",
    "TOTAL PASIVO", "TOTAL PATRIMONIO", "TOTAL PASIVO + PATRIMONIO",
    "Ventas brutas (total)",
}

SKIP_LABELS = {
    'Ventas y unidades vienen de la hoja "Ventas por Producto 2026" (texto verde = vínculo a otra hoja).',
    "Control: Activo − (Pasivo + Patrimonio)",
}


def fmt(v, is_header=False):
    if v is None:
        return ""
    if is_header:
        v = str(v).replace(" (Dec 31)", "")
        return v
    if isinstance(v, (int, float)):
        iv = round(v)
        if iv < 0:
            return f"-${abs(iv):,}"
        return f"${iv:,}"
    return str(v)


def render_sheet(sheet_name, title, filename, label_width=0.26):
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            continue
        if row[0] in SKIP_LABELS:
            continue
        rows.append(row)

    header = rows[0]
    data_rows = rows[1:]
    n_rows = len(data_rows)
    n_cols = len(header)

    fig_w = 18.5
    row_h = 0.30          # inches per data row
    header_h = 0.42       # inches
    title_h = 0.75        # inches reserved at top for the title
    bottom_pad = 0.15     # inches of blank margin at the very bottom

    content_h = header_h + row_h * n_rows
    fig_h = title_h + content_h + bottom_pad

    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=200)
    ax.set_axis_off()
    ax.set_xlim(0, fig_w)
    ax.set_ylim(0, fig_h)

    ax.text(0.15, fig_h - title_h / 2, title, fontsize=15, fontweight="bold", color=BROWN,
             ha="left", va="center", family="sans-serif")

    top = fig_h - title_h
    data_col_w = (fig_w - label_width * fig_w) / (n_cols - 1)
    label_w_in = label_width * fig_w

    # header row
    ax.add_patch(plt.Rectangle((0, top - header_h), fig_w, header_h, color=BROWN, zorder=1))
    ax.text(0.15, top - header_h / 2, str(header[0]), fontsize=9.5, fontweight="bold",
            color=WHITE, ha="left", va="center", zorder=2)
    for c in range(1, n_cols):
        x = label_w_in + (c - 1) * data_col_w
        ax.text(x + data_col_w - 0.08, top - header_h / 2, fmt(header[c], True),
                fontsize=9.5, fontweight="bold", color=WHITE, ha="right", va="center", zorder=2)

    y = top - header_h
    for r, row in enumerate(data_rows):
        label = row[0]
        is_group = label in GROUP_LABELS
        is_total = label in TOTAL_LABELS
        y -= row_h

        if is_group:
            ax.add_patch(plt.Rectangle((0, y), fig_w, row_h, color=LIGHT, zorder=1))
            ax.text(0.15, y + row_h / 2, str(label), fontsize=9, fontweight="bold",
                    color=BROWN, ha="left", va="center", zorder=2)
            continue

        if r % 2 == 0:
            ax.add_patch(plt.Rectangle((0, y), fig_w, row_h, color="#FBF6EC", zorder=1))
        if is_total:
            ax.add_patch(plt.Rectangle((0, y), fig_w, row_h, color="#EADFC8", zorder=1))

        weight = "bold" if is_total else "normal"
        ax.text(0.15, y + row_h / 2, str(label), fontsize=8.5, fontweight=weight,
                color=DARK, ha="left", va="center", zorder=2)

        for c in range(1, n_cols):
            val = row[c]
            x = label_w_in + (c - 1) * data_col_w
            color = DARK
            if isinstance(val, (int, float)) and val < 0:
                color = RED
            txt = fmt(val)
            ax.text(x + data_col_w - 0.08, y + row_h / 2, txt, fontsize=8.2,
                    fontweight=weight, color=color, ha="right", va="center", zorder=2)

    # outer border
    total_h = top - y
    ax.add_patch(plt.Rectangle((0, y), fig_w, total_h, fill=False, edgecolor=BROWN, linewidth=1.2, zorder=3))
    # column separator between label and data
    ax.plot([label_w_in, label_w_in], [y, top], color=BROWN, linewidth=0.8, zorder=3)
    # separator before Annual column
    x_annual = label_w_in + (n_cols - 2) * data_col_w
    ax.plot([x_annual, x_annual], [y, top], color=BROWN, linewidth=0.8, zorder=3)

    plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
    out_path = f"{OUT_DIR}/{filename}"
    fig.savefig(out_path, facecolor=WHITE)
    plt.close(fig)
    print("Saved", out_path, n_rows, "rows x", n_cols, "cols")


render_sheet(
    "Estado de Resultados 2026",
    "Estado de Resultados — Döner Street, Año 1 (2026, Realistic)",
    "estado-resultados-2026.png",
    label_width=0.22,
)
render_sheet(
    "Flujo de Efectivo 2026",
    "Flujo de Efectivo — Döner Street, Año 1 (2026, Realistic)",
    "flujo-efectivo-2026.png",
    label_width=0.26,
)
render_sheet(
    "Balance General 2026",
    "Balance General — Döner Street, Año 1 (2026, Realistic)",
    "balance-general-2026.png",
    label_width=0.26,
)
